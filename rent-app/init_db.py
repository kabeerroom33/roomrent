"""
Initialize the SQLite database and import all data from Excel file.
Run this once to set up the app.
"""
import sqlite3
import msoffcrypto
import io
import openpyxl
import os

DB_PATH = os.path.join(os.path.dirname(__file__), 'data', 'rent.db')
EXCEL_PATH = os.path.join(os.path.dirname(__file__), '..', 'ROOM RENT JAN TO DEC 2025.xlsx')
PASSWORD = '9745437665'

MONTHS = ['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec']

def create_tables(conn):
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS clients (
            id          INTEGER PRIMARY KEY,
            name        TEXT NOT NULL,
            mobile      TEXT DEFAULT '',
            room_no     TEXT DEFAULT '',
            monthly_rent REAL DEFAULT 525,
            active      INTEGER DEFAULT 1,
            is_hidden   INTEGER DEFAULT 0,
            notes       TEXT DEFAULT '',
            created_at  TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS payments (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            client_id   INTEGER NOT NULL,
            year        INTEGER NOT NULL,
            month       TEXT NOT NULL,
            amount_due  REAL DEFAULT 0,
            amount_paid REAL DEFAULT 0,
            balance     REAL DEFAULT 0,
            paid_date   TEXT,
            receipt_no  TEXT,
            notes       TEXT DEFAULT '',
            payment_method TEXT DEFAULT 'cash',
            FOREIGN KEY(client_id) REFERENCES clients(id)
        );

        CREATE TABLE IF NOT EXISTS old_balances (
            client_id   INTEGER PRIMARY KEY,
            year        INTEGER NOT NULL,
            amount      REAL DEFAULT 0,
            FOREIGN KEY(client_id) REFERENCES clients(id)
        );
    """)
    client_columns = {row[1] for row in conn.execute('PRAGMA table_info(clients)')}
    if 'is_hidden' not in client_columns:
        conn.execute('ALTER TABLE clients ADD COLUMN is_hidden INTEGER DEFAULT 0')
    payment_columns = {row[1] for row in conn.execute('PRAGMA table_info(payments)')}
    if 'payment_method' not in payment_columns:
        conn.execute("ALTER TABLE payments ADD COLUMN payment_method TEXT DEFAULT 'cash'")
    conn.commit()

def import_excel(conn):
    # Decrypt and load Excel
    with open(EXCEL_PATH, 'rb') as f:
        office = msoffcrypto.OfficeFile(f)
        office.load_key(password=PASSWORD)
        decrypted = io.BytesIO()
        office.decrypt(decrypted)

    decrypted.seek(0)
    wb = openpyxl.load_workbook(decrypted, data_only=True)
    rent_ws = wb['Rent']
    paid_ws = wb['Paid']

    # Build lookup: row_num -> paid amounts
    paid_lookup = {}
    for row in paid_ws.iter_rows(min_row=5, max_row=paid_ws.max_row, values_only=True):
        no = row[1]
        if no is None:
            continue
        old_bal_paid = row[3] or 0
        monthly_paid = list(row[4:16])  # Jan-Dec
        paid_lookup[no] = {'old_bal': old_bal_paid, 'months': monthly_paid}

    # Import clients and payments from Rent sheet
    for row in rent_ws.iter_rows(min_row=5, max_row=rent_ws.max_row, values_only=True):
        no = row[1]
        name = row[2]
        if no is None or name is None or str(name).startswith('=') or str(name).strip() in ('', 'A'):
            continue

        name = str(name).strip()
        old_balance_due = row[3] or 0
        monthly_dues = list(row[4:16])  # Jan-Dec rent amounts

        # Insert client
        conn.execute("""
            INSERT OR REPLACE INTO clients (id, name, room_no, monthly_rent, active)
            VALUES (?, ?, 'Room 33', 525, 1)
        """, (no, name))

        # Old balance
        paid_old = paid_lookup.get(no, {}).get('old_bal', 0)
        old_balance_net = old_balance_due - paid_old
        if old_balance_due != 0 or paid_old != 0:
            conn.execute("""
                INSERT OR REPLACE INTO old_balances (client_id, year, amount)
                VALUES (?, 2025, ?)
            """, (no, old_balance_net))

        # Monthly payments
        paid_months = paid_lookup.get(no, {}).get('months', [None]*12)
        for i, month in enumerate(MONTHS):
            due = monthly_dues[i] or 0
            paid = paid_months[i] or 0
            if due == 0 and paid == 0:
                continue
            balance = due - paid
            conn.execute("""
                INSERT OR REPLACE INTO payments
                    (client_id, year, month, amount_due, amount_paid, balance)
                VALUES (?, 2025, ?, ?, ?, ?)
            """, (no, month, due, paid, balance))

    conn.commit()
    print(f"✅ Import complete! Clients: {conn.execute('SELECT COUNT(*) FROM clients').fetchone()[0]}")
    print(f"   Payments: {conn.execute('SELECT COUNT(*) FROM payments').fetchone()[0]}")

if __name__ == '__main__':
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    create_tables(conn)
    import_excel(conn)
    conn.close()
    print("Database ready at:", DB_PATH)
