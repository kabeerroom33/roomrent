"""
Room Rent Management App — Room 33
Flask backend: SQLite locally, PostgreSQL (Supabase) in production.
"""
from flask import (Flask, abort, render_template, request, redirect,
                   url_for, send_file, flash, session, jsonify)
import os, io, hmac, secrets
from datetime import datetime
from xml.sax.saxutils import escape
from werkzeug.security import generate_password_hash, check_password_hash
from dotenv import load_dotenv

load_dotenv()

from app_config import DB_PATH, ROOM_NAME
from db import get_db, init_schema

# ─── App setup ─────────────────────────────────────────────────────────────────
app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY') or secrets.token_hex(32)
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    SESSION_COOKIE_SECURE=os.environ.get('RENDER', '') != '',
)

MONTHS = ['jan','feb','mar','apr','may','jun','jul','aug','sep','oct','nov','dec']
MONTH_NAMES = {
    'jan':'January','feb':'February','mar':'March','apr':'April',
    'may':'May','jun':'June','jul':'July','aug':'August',
    'sep':'September','oct':'October','nov':'November','dec':'December'
}
ADMIN_PASSWORD_HASH = os.environ.get(
    'ADMIN_PASSWORD_HASH',
    generate_password_hash('admin1234')   # default — override in .env
)

# ─── CSRF + Auth ───────────────────────────────────────────────────────────────
@app.before_request
def protect():
    if request.method == 'POST' and request.endpoint != 'login':
        token = request.form.get('csrf_token', '')
        expected = session.get('csrf_token', '')
        if not expected or not hmac.compare_digest(token, expected):
            abort(400, 'Security token mismatch. Reload the page and try again.')
    if request.endpoint in ('login', 'static', 'health'):
        return
    if not session.get('authenticated'):
        return redirect(url_for('login', next=request.path))

@app.context_processor
def inject_globals():
    token = session.get('csrf_token')
    if not token:
        token = secrets.token_urlsafe(32)
        session['csrf_token'] = token
    return {'csrf_token': token, 'room_name': ROOM_NAME,
            'months': MONTHS, 'month_names': MONTH_NAMES}

# ─── Auth routes ───────────────────────────────────────────────────────────────
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        pw = request.form.get('password', '')
        if check_password_hash(ADMIN_PASSWORD_HASH, pw):
            session['authenticated'] = True
            session.permanent = True
            return redirect(request.args.get('next') or url_for('dashboard'))
        flash('Wrong password', 'error')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/health')
def health():
    return 'OK', 200

# ─── Helpers ───────────────────────────────────────────────────────────────────
def _month_order_sql(backend='sq'):
    """Inline CASE for ordering months."""
    return ("CASE month WHEN 'jan' THEN 1 WHEN 'feb' THEN 2 WHEN 'mar' THEN 3 "
            "WHEN 'apr' THEN 4 WHEN 'may' THEN 5 WHEN 'jun' THEN 6 "
            "WHEN 'jul' THEN 7 WHEN 'aug' THEN 8 WHEN 'sep' THEN 9 "
            "WHEN 'oct' THEN 10 WHEN 'nov' THEN 11 WHEN 'dec' THEN 12 END")

def get_available_years():
    years = {datetime.now().year}
    with get_db() as conn:
        conn.execute('SELECT year FROM payments UNION SELECT year FROM old_balances')
        years.update(int(row['year']) for row in conn.fetchall())
    return sorted(years, reverse=True)

def get_client_full(client_id, year=None):
    if year is None:
        year = datetime.now().year
    with get_db() as conn:
        conn.execute("SELECT * FROM clients WHERE id=?", (client_id,))
        client = conn.fetchone()
        if not client:
            return None
        conn.execute(
            "SELECT amount FROM old_balances WHERE client_id=? AND year=?",
            (client_id, year))
        old_bal = conn.fetchone()
        conn.execute(
            f"SELECT * FROM payments WHERE client_id=? AND year=? ORDER BY {_month_order_sql()}",
            (client_id, year))
        payments = conn.fetchall()

    payment_map   = {p['month']: dict(p) for p in payments}
    total_due     = sum(float(p['amount_due'])  for p in payments)
    total_paid    = sum(float(p['amount_paid']) for p in payments)
    total_balance = sum(float(p['balance'])     for p in payments)
    ob = float(old_bal['amount']) if old_bal else 0
    total_balance += ob

    return {
        'client': dict(client),
        'old_balance': ob,
        'payments': payment_map,
        'total_due': total_due,
        'total_paid': total_paid,
        'total_balance': total_balance,
    }

# ─── Dashboard ─────────────────────────────────────────────────────────────────
@app.route('/')
def dashboard():
    year = int(request.args.get('year', datetime.now().year))
    years = get_available_years()
    show_hidden = request.args.get('show_hidden') == '1'

    with get_db() as conn:
        if show_hidden:
            conn.execute("SELECT * FROM clients WHERE active=1 ORDER BY id")
        else:
            conn.execute("SELECT * FROM clients WHERE active=1 AND is_hidden=0 ORDER BY id")
        clients = conn.fetchall()

    summary = []
    total_due_all = total_paid_all = total_balance_all = 0

    for c in clients:
        with get_db() as conn:
            conn.execute(
                "SELECT amount FROM old_balances WHERE client_id=? AND year=?",
                (c['id'], year))
            old_bal = conn.fetchone()
            conn.execute(
                "SELECT COALESCE(SUM(amount_due),0) as due, "
                "       COALESCE(SUM(amount_paid),0) as paid, "
                "       COALESCE(SUM(balance),0) as bal "
                "FROM payments WHERE client_id=? AND year=?",
                (c['id'], year))
            stats = conn.fetchone()

        due  = float(stats['due'])
        paid = float(stats['paid'])
        bal  = float(stats['bal']) + (float(old_bal['amount']) if old_bal else 0)
        total_due_all     += due
        total_paid_all    += paid
        total_balance_all += bal
        summary.append({
            'id': c['id'], 'name': c['name'],
            'mobile': c['mobile'], 'room_no': c['room_no'],
            'is_hidden': c['is_hidden'],
            'due': due, 'paid': paid, 'balance': bal
        })

    return render_template('dashboard.html',
        summary=summary, year=year,
        total_due=total_due_all, total_paid=total_paid_all,
        total_balance=total_balance_all,
        show_hidden=show_hidden,
        years=years,
    )

# ─── Client detail ─────────────────────────────────────────────────────────────
@app.route('/client/<int:client_id>')
def client_detail(client_id):
    year = int(request.args.get('year', datetime.now().year))
    years = get_available_years()
    data = get_client_full(client_id, year)
    if not data:
        flash('Client not found', 'error')
        return redirect(url_for('dashboard'))
    return render_template('client_detail.html', data=data, year=year, years=years)

# ─── Add client ────────────────────────────────────────────────────────────────
@app.route('/client/add', methods=['GET', 'POST'])
def add_client():
    if request.method == 'POST':
        with get_db() as conn:
            conn.execute("SELECT COALESCE(MAX(id),0)+1 FROM clients")
            new_id = conn.fetchone()[0]
            conn.execute("""
                INSERT INTO clients (id, name, mobile, room_no, monthly_rent, notes)
                VALUES (?,?,?,?,?,?)
            """, (
                new_id,
                request.form['name'].strip().upper(),
                request.form.get('mobile','').strip(),
                request.form.get('room_no', ROOM_NAME).strip() or ROOM_NAME,
                float(request.form.get('monthly_rent', 525)),
                request.form.get('notes','').strip()
            ))
        flash(f'Client added — ID #{new_id}', 'success')
        return redirect(url_for('client_detail', client_id=new_id))
    return render_template('add_client.html', default_room=ROOM_NAME)

# ─── Edit client ───────────────────────────────────────────────────────────────
@app.route('/client/<int:client_id>/edit', methods=['GET', 'POST'])
def edit_client(client_id):
    with get_db() as conn:
        conn.execute("SELECT * FROM clients WHERE id=?", (client_id,))
        client = conn.fetchone()
    if not client:
        flash('Client not found', 'error')
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        with get_db() as conn:
            conn.execute("""
                UPDATE clients SET name=?,mobile=?,room_no=?,monthly_rent=?,notes=?
                WHERE id=?
            """, (
                request.form['name'].strip().upper(),
                request.form.get('mobile','').strip(),
                request.form.get('room_no', ROOM_NAME).strip() or ROOM_NAME,
                float(request.form.get('monthly_rent', 525)),
                request.form.get('notes','').strip(),
                client_id
            ))
        flash('Client updated', 'success')
        return redirect(url_for('client_detail', client_id=client_id))
    return render_template('edit_client.html', client=dict(client))

# ─── Hide / Restore / Delete ──────────────────────────────────────────────────
@app.route('/client/<int:client_id>/hide', methods=['POST'])
def hide_client(client_id):
    with get_db() as conn:
        conn.execute("UPDATE clients SET is_hidden=1 WHERE id=?", (client_id,))
    flash('Client hidden from list', 'success')
    return redirect(request.referrer or url_for('dashboard'))

@app.route('/client/<int:client_id>/restore', methods=['POST'])
def restore_client(client_id):
    with get_db() as conn:
        conn.execute("UPDATE clients SET is_hidden=0 WHERE id=?", (client_id,))
    flash('Client restored', 'success')
    return redirect(request.referrer or url_for('dashboard'))

@app.route('/client/<int:client_id>/delete', methods=['POST'])
def delete_client(client_id):
    with get_db() as conn:
        conn.execute("DELETE FROM payments WHERE client_id=?", (client_id,))
        conn.execute("DELETE FROM old_balances WHERE client_id=?", (client_id,))
        conn.execute("DELETE FROM clients WHERE id=?", (client_id,))
    flash('Client permanently deleted', 'success')
    return redirect(url_for('dashboard'))

# ─── Record payment ────────────────────────────────────────────────────────────
@app.route('/payment/record', methods=['POST'])
def record_payment():
    client_id      = int(request.form['client_id'])
    year           = int(request.form['year'])
    month          = request.form['month']
    paid           = float(request.form['amount_paid'])
    due            = float(request.form.get('amount_due', 525))
    notes          = request.form.get('notes', '')
    paid_date      = request.form.get('paid_date', '') or datetime.now().strftime('%Y-%m-%d')
    payment_method = request.form.get('payment_method', 'cash')

    with get_db() as conn:
        conn.execute(
            "SELECT id, amount_paid, amount_due FROM payments "
            "WHERE client_id=? AND year=? AND month=?",
            (client_id, year, month))
        existing = conn.fetchone()

        if existing:
            new_paid = float(existing['amount_paid']) + paid
            new_due  = float(existing['amount_due'])
            balance  = new_due - new_paid
            conn.execute("""
                UPDATE payments
                SET amount_paid=?, balance=?, paid_date=?, payment_method=?, notes=?
                WHERE id=?
            """, (new_paid, balance, paid_date, payment_method, notes, existing['id']))
        else:
            balance = due - paid
            conn.execute("""
                INSERT INTO payments
                  (client_id,year,month,amount_due,amount_paid,balance,paid_date,payment_method,notes)
                VALUES (?,?,?,?,?,?,?,?,?)
            """, (client_id, year, month, due, paid, balance,
                  paid_date, payment_method, notes))

    flash(f'Payment of AED {paid:.0f} recorded for {MONTH_NAMES[month]} {year}', 'success')
    return redirect(url_for('client_detail', client_id=client_id, year=year))

@app.route('/payment/<int:payment_id>/edit', methods=['POST'])
def edit_payment(payment_id):
    due = float(request.form.get('amount_due', 0))
    paid = float(request.form.get('amount_paid', 0))
    paid_date = request.form.get('paid_date', '') or None
    payment_method = request.form.get('payment_method', 'cash')
    notes = request.form.get('notes', '').strip()

    with get_db() as conn:
        conn.execute('SELECT client_id, year, month FROM payments WHERE id=?', (payment_id,))
        payment = conn.fetchone()
        if not payment:
            flash('Payment entry not found', 'error')
            return redirect(url_for('dashboard'))
        conn.execute(
            "UPDATE payments SET amount_due=?, amount_paid=?, balance=?, paid_date=?, payment_method=?, notes=? WHERE id=?",
            (due, paid, due - paid, paid_date, payment_method, notes, payment_id))

    flash(f'{MONTH_NAMES[payment["month"]]} {payment["year"]} payment updated', 'success')
    return redirect(url_for('client_detail', client_id=payment['client_id'], year=payment['year']))

@app.route('/payment/<int:payment_id>/delete', methods=['POST'])
def delete_payment(payment_id):
    with get_db() as conn:
        conn.execute('SELECT client_id, year, month FROM payments WHERE id=?', (payment_id,))
        payment = conn.fetchone()
        if not payment:
            flash('Payment entry not found', 'error')
            return redirect(url_for('dashboard'))
        conn.execute('DELETE FROM payments WHERE id=?', (payment_id,))

    flash(f'{MONTH_NAMES[payment["month"]]} {payment["year"]} payment entry deleted', 'success')
    return redirect(url_for('client_detail', client_id=payment['client_id'], year=payment['year']))

# ─── Search ────────────────────────────────────────────────────────────────────
@app.route('/search')
def search():
    q = request.args.get('q', '').strip()
    results = []
    if q:
        with get_db() as conn:
            if q.isdigit():
                conn.execute(
                    "SELECT * FROM clients WHERE (id=? OR mobile LIKE ?) AND active=1 ORDER BY id",
                    (int(q), f'%{q}%'))
            else:
                conn.execute(
                    "SELECT * FROM clients WHERE name LIKE ? AND active=1 ORDER BY id",
                    (f'%{q.upper()}%',))
            results = [dict(r) for r in conn.fetchall()]
    return render_template('search.html', results=results, query=q)

# ─── Monthly payments report ───────────────────────────────────────────────────
@app.route('/reports/monthly-payments')
def monthly_payments_report():
    year = int(request.args.get('year', datetime.now().year))
    years = get_available_years()

    with get_db() as conn:
        conn.execute("SELECT * FROM clients WHERE active=1 AND is_hidden=0 ORDER BY id")
        clients = conn.fetchall()

    monthly_totals = {month: 0.0 for month in MONTHS}
    report_rows = []

    for client in clients:
        monthly_amounts = {}
        total = 0.0
        for month in MONTHS:
            with get_db() as conn:
                conn.execute(
                    "SELECT COALESCE(SUM(amount_paid), 0) AS paid FROM payments WHERE client_id=? AND year=? AND month=?",
                    (client['id'], year, month)
                )
                payment = conn.fetchone()
            amount = float(payment['paid']) if payment else 0.0
            monthly_amounts[month] = amount
            total += amount
            monthly_totals[month] += amount

        report_rows.append({
            'id': client['id'],
            'name': client['name'],
            'room_no': client['room_no'],
            'mobile': client['mobile'],
            'months': monthly_amounts,
            'total': total,
        })

    report_rows.sort(key=lambda row: row['total'], reverse=True)
    total_collected = sum(monthly_totals.values())

    return render_template(
        'monthly_payments_report.html',
        rows=report_rows,
        year=year,
        years=years,
        months=MONTHS,
        monthly_totals=monthly_totals,
        total_collected=total_collected,
    )

# ─── Dues report ───────────────────────────────────────────────────────────────
@app.route('/dues')
def dues_report():
    year = int(request.args.get('year', datetime.now().year))
    years = get_available_years()
    with get_db() as conn:
        conn.execute("SELECT * FROM clients WHERE active=1 AND is_hidden=0 ORDER BY id")
        clients = conn.fetchall()
    dues = []
    for c in clients:
        with get_db() as conn:
            conn.execute(
                "SELECT amount FROM old_balances WHERE client_id=? AND year=?",
                (c['id'], year))
            old_bal = conn.fetchone()
            conn.execute(
                "SELECT COALESCE(SUM(balance),0) as bal FROM payments WHERE client_id=? AND year=?",
                (c['id'], year))
            stats = conn.fetchone()
        bal = float(stats['bal']) + (float(old_bal['amount']) if old_bal else 0)
        if bal > 0:
            dues.append({'id': c['id'], 'name': c['name'],
                         'mobile': c['mobile'], 'balance': bal})
    dues.sort(key=lambda x: x['balance'], reverse=True)
    return render_template('dues.html', dues=dues, year=year,
                           total=sum(d['balance'] for d in dues), years=years)

# ─── PDF Receipt ───────────────────────────────────────────────────────────────
@app.route('/client/<int:client_id>/receipt/<int:year>/<month>')
def generate_receipt(client_id, year, month):
    from pdf_generator import create_receipt_pdf
    data = get_client_full(client_id, year)
    if not data:
        flash('Client not found', 'error')
        return redirect(url_for('dashboard'))
    buf = create_receipt_pdf(data, year, month, MONTH_NAMES, ROOM_NAME)
    fname = (f"Receipt_{data['client']['name'].replace(' ','_')}"
             f"_{MONTH_NAMES[month]}_{year}.pdf")
    return send_file(io.BytesIO(buf), mimetype='application/pdf',
                     as_attachment=True, download_name=fname)

@app.route('/client/<int:client_id>/receipt-all/<int:year>')
def generate_full_statement(client_id, year):
    from pdf_generator import create_full_statement_pdf
    data = get_client_full(client_id, year)
    if not data:
        flash('Client not found', 'error')
        return redirect(url_for('dashboard'))
    buf = create_full_statement_pdf(data, year, MONTHS, MONTH_NAMES, ROOM_NAME)
    fname = f"Statement_{data['client']['name'].replace(' ','_')}_{year}.pdf"
    return send_file(io.BytesIO(buf), mimetype='application/pdf',
                     as_attachment=True, download_name=fname)

# ─── Import Excel data ─────────────────────────────────────────────────────────
@app.route('/admin/import', methods=['GET', 'POST'])
def import_excel():
    """Import or re-import data from Excel files."""
    message = None
    years = get_available_years()
    if request.method == 'POST':
        year_value = request.form.get('year', '').strip()
        admin_password = request.form.get('admin_password', '')
        workbook_password = request.form.get('workbook_password', '').strip()
        if not year_value.isdigit() or int(year_value) not in years:
            message = ('error', 'Choose a year before importing.')
        elif not check_password_hash(ADMIN_PASSWORD_HASH, admin_password):
            message = ('error', 'The admin password is incorrect.')
        elif not workbook_password:
            message = ('error', 'Enter the Excel workbook password before importing.')
        else:
            year = int(year_value)
            try:
                count = _do_import(year, workbook_password)
                message = ('success', f'✅ Imported {count} payment records for {year}')
            except Exception as e:
                message = ('error', f'Error: {e}')
    return render_template('import.html', message=message, years=years)

def _do_import(year, password):
    import msoffcrypto, openpyxl
    months = ['jan','feb','mar','apr','may','jun','jul','aug','sep','oct','nov','dec']
    fname = os.path.join(os.path.dirname(__file__), '..', f'ROOM RENT JAN TO DEC {year}.xlsx')
    if not os.path.exists(fname):
        raise FileNotFoundError(f'File not found: {fname}')

    with open(fname, 'rb') as f:
        office = msoffcrypto.OfficeFile(f)
        office.load_key(password=password)
        dec = io.BytesIO()
        office.decrypt(dec)
    dec.seek(0)
    wb = openpyxl.load_workbook(dec, data_only=True)
    rent_ws = wb['Rent']
    paid_ws = wb['Paid']

    paid_lookup = {}
    for row in paid_ws.iter_rows(min_row=5, max_row=paid_ws.max_row, values_only=True):
        no = row[1]
        if no is None: continue
        paid_lookup[no] = {'old_bal': row[3] or 0, 'months': list(row[4:16])}

    # Clear existing data for the year
    with get_db() as conn:
        conn.execute("DELETE FROM payments WHERE year=?", (year,))
        conn.execute("DELETE FROM old_balances WHERE year=?", (year,))

    count = 0
    for row in rent_ws.iter_rows(min_row=5, max_row=rent_ws.max_row, values_only=True):
        no   = row[1]
        name = row[2]
        if no is None or name is None or str(name).startswith('='): continue
        name = str(name).strip()
        if name in ('', 'A'): continue

        old_due  = row[3] or 0
        mon_dues = list(row[4:16])

        # Upsert client
        with get_db() as conn:
            conn.execute("SELECT id, name FROM clients WHERE id=?", (no,))
            existing = conn.fetchone()
        if not existing:
            with get_db() as conn:
                conn.execute(
                    "INSERT INTO clients (id, name, monthly_rent, room_no) VALUES (?,?,525,?)",
                    (no, name, ROOM_NAME))
        elif existing['name'].strip() in ('A', ''):
            with get_db() as conn:
                conn.execute("UPDATE clients SET name=? WHERE id=?", (name, no))

        # Old balance
        paid_old = paid_lookup.get(no, {}).get('old_bal', 0)
        ob_net   = old_due - paid_old
        if old_due != 0 or paid_old != 0:
            with get_db() as conn:
                conn.execute(
                    "INSERT INTO old_balances (client_id, year, amount) VALUES (?,?,?) "
                    "ON CONFLICT(client_id, year) DO UPDATE SET amount=EXCLUDED.amount"
                    if os.environ.get('DATABASE_URL') else
                    "INSERT OR REPLACE INTO old_balances (client_id, year, amount) VALUES (?,?,?)",
                    (no, year, ob_net))

        # Monthly payments
        paid_months = paid_lookup.get(no, {}).get('months', [None]*12)
        for i, month in enumerate(months):
            due  = mon_dues[i] or 0
            paid = paid_months[i] or 0
            if due == 0 and paid == 0: continue
            with get_db() as conn:
                conn.execute("""
                    INSERT INTO payments (client_id,year,month,amount_due,amount_paid,balance)
                    VALUES (?,?,?,?,?,?)
                """, (no, year, month, due, paid, due - paid))
            count += 1

    return count

# ─── Backup / export ───────────────────────────────────────────────────────────
@app.route('/admin/backup')
def backup():
    return redirect(url_for('dashboard'))

def _collect_export_data():
    with get_db() as conn:
        conn.execute('SELECT * FROM clients ORDER BY id')
        clients = [dict(row) for row in conn.fetchall()]
        conn.execute(
            "SELECT p.*, c.name FROM payments p JOIN clients c ON c.id=p.client_id "
            "ORDER BY p.year, c.id, CASE p.month "
            "WHEN 'jan' THEN 1 WHEN 'feb' THEN 2 WHEN 'mar' THEN 3 "
            "WHEN 'apr' THEN 4 WHEN 'may' THEN 5 WHEN 'jun' THEN 6 "
            "WHEN 'jul' THEN 7 WHEN 'aug' THEN 8 WHEN 'sep' THEN 9 "
            "WHEN 'oct' THEN 10 WHEN 'nov' THEN 11 WHEN 'dec' THEN 12 END")
        payments = [dict(row) for row in conn.fetchall()]
        conn.execute('SELECT * FROM old_balances ORDER BY year, client_id')
        old_balances = [dict(row) for row in conn.fetchall()]

    years = {datetime.now().year}
    years.update(int(payment['year']) for payment in payments)
    years.update(int(balance['year']) for balance in old_balances)
    return clients, payments, old_balances, sorted(years)

@app.route('/admin/export/excel')
def export_excel():
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    clients, payments, old_balances, years = _collect_export_data()
    template_path = os.path.join(os.path.dirname(__file__), '..', 'sample.xlsx')
    if os.path.exists(template_path):
        workbook = load_workbook(template_path)
        report = workbook.active
        report.delete_rows(1, report.max_row)
    else:
        workbook = Workbook()
        report = workbook.active
    report.title = 'Rent Report'
    headers = ('YEAR', 'CLIENTS', 'MONTH', 'DUE', 'PAID', 'BALANCE',
               'TOTAL OUTSTANDING BALANCE', 'NOTES')
    report.append(headers)

    month_index = {month: index for index, month in enumerate(MONTHS)}
    payments_by_key = {
        (int(payment['client_id']), int(payment['year']), payment['month']): payment
        for payment in payments
    }
    balances_by_key = {
        (int(balance['client_id']), int(balance['year'])): float(balance['amount'] or 0)
        for balance in old_balances
    }
    header_fill = PatternFill('solid', fgColor='1A3C6E')
    header_font = Font(color='FFFFFF', bold=True)
    border_side = Side(style='thin', color='D9E2F2')
    currency_format = '#,##0.00;[Red]-#,##0.00;–'

    for year in years:
        for client in clients:
            client_id = int(client['id'])
            annual_payments = [
                payments_by_key.get((client_id, year, month), {}) for month in MONTHS
            ]
            total_outstanding = balances_by_key.get((client_id, year), 0) + sum(
                float(payment.get('balance') or 0) for payment in annual_payments)
            first_row = report.max_row + 1
            for index, month in enumerate(MONTHS):
                payment = annual_payments[index]
                report.append((
                    year if index == 0 else None,
                    client['name'] if index == 0 else None,
                    MONTH_NAMES[month].upper(),
                    float(payment.get('amount_due') or 0),
                    float(payment.get('amount_paid') or 0),
                    float(payment.get('balance') or 0),
                    total_outstanding if index == 0 else None,
                    payment.get('notes') or '',
                ))
            last_row = report.max_row
            for column in (1, 2, 7):
                report.merge_cells(start_row=first_row, start_column=column,
                                   end_row=last_row, end_column=column)

    for cell in report[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(bottom=Side(style='medium', color='F59E0B'))
    report.row_dimensions[1].height = 32
    for row in report.iter_rows(min_row=2, min_col=1, max_col=8):
        for cell in row:
            cell.border = Border(bottom=border_side)
            cell.alignment = Alignment(vertical='top', wrap_text=(cell.column in (2, 8)))
            if cell.column in (4, 5, 6, 7):
                cell.number_format = currency_format
    for column, width in {'A': 10, 'B': 34, 'C': 16, 'D': 14, 'E': 14,
                          'F': 14, 'G': 23, 'H': 42}.items():
        report.column_dimensions[column].width = width
    report.freeze_panes = 'A2'
    report.sheet_view.showGridLines = False

    client_sheet = workbook.create_sheet('Clients')
    client_sheet.append(('ID', 'Name', 'Mobile', 'Room', 'Monthly Rent', 'Active', 'Hidden', 'Client Notes'))
    for client in clients:
        client_sheet.append((client['id'], client['name'], client['mobile'], client['room_no'],
                             client['monthly_rent'], client['active'], client['is_hidden'], client['notes']))
    payment_sheet = workbook.create_sheet('Transactions')
    payment_sheet.append(('Payment ID', 'Client ID', 'Client', 'Year', 'Month', 'Due', 'Paid',
                          'Balance', 'Paid Date', 'Method', 'Notes'))
    for payment in payments:
        payment_sheet.append((payment['id'], payment['client_id'], payment['name'], payment['year'],
                              MONTH_NAMES[payment['month']], payment['amount_due'], payment['amount_paid'],
                              payment['balance'], payment['paid_date'], payment['payment_method'], payment['notes']))
    balance_sheet = workbook.create_sheet('Opening Balances')
    balance_sheet.append(('Client ID', 'Client', 'Year', 'Opening Balance'))
    client_names = {int(client['id']): client['name'] for client in clients}
    for balance in old_balances:
        balance_sheet.append((balance['client_id'], client_names.get(int(balance['client_id']), ''),
                              balance['year'], balance['amount']))

    for sheet in (client_sheet, payment_sheet, balance_sheet):
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        for cells in sheet.columns:
            width = min(max(max(len(str(cell.value or '')) for cell in cells) + 2, 12), 42)
            sheet.column_dimensions[cells[0].column_letter].width = width

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='sample.xlsx')

@app.route('/admin/export/pdf')
def export_pdf():
    from pdf_generator import create_export_pdf
    clients, payments, old_balances, years = _collect_export_data()
    report = create_export_pdf(clients, payments, old_balances, years, MONTHS, MONTH_NAMES, ROOM_NAME)
    return send_file(io.BytesIO(report), mimetype='application/pdf', as_attachment=True,
                     download_name='rent-data-report.pdf')

# ─── Startup ───────────────────────────────────────────────────────────────────
def create_app():
    os.makedirs(os.path.join(os.path.dirname(__file__), 'data'), exist_ok=True)
    init_schema()
    return app

if __name__ == '__main__':
    create_app()
    print(f"🏠 Room 33 — Rent Manager  →  http://localhost:8080")
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 8080)), debug=False)
